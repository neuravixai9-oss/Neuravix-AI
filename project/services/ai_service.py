"""
Ядро общения с Google AI Studio (Gemini).
Версия 3.2 — ускорен стриминг, исправлена генерация изображений,
добавлены таймауты и надёжный retry.
"""

import asyncio
import io
import logging
import zipfile

from google import genai
from google.genai import types
from google.genai import errors as genai_errors

from config import (
    GEMINI_API_KEY, IMAGE_GEN_MODELS, UTILITY_MODEL,
    CHAT_FALLBACK_MODEL, SUPPORT_USERNAME, SUPER_OWNER_ID, SUPER_OWNER_USERNAME,
)

logger = logging.getLogger("ai_service")

_client: genai.Client | None = None

STREAM_TIMEOUT = 45.0   # секунд — таймаут всего стрима
EDIT_INTERVAL = 0.5     # секунд между правками сообщения


def get_client() -> genai.Client:
    global _client
    if _client is None:
        if not GEMINI_API_KEY:
            raise RuntimeError("GEMINI_API_KEY не настроен")
        _client = genai.Client(api_key=GEMINI_API_KEY)
    return _client


BASE_SYSTEM_PROMPT = (
    "Ты — Neuravix AI, умный, дружелюбный и предельно полезный универсальный "
    "ИИ-ассистент внутри Telegram-бота. Пользователь общается с тобой в ОДНОМ "
    "чате и ожидает, что ты сам поймёшь задачу и выполнишь её — без выбора "
    "отдельных режимов или кнопок.\n\n"
    "Ты умеешь прямо в этом чате:\n"
    "• Отвечать на любые вопросы и поддерживать беседу с памятью контекста диалога.\n"
    "• Переводить текст на любой язык — просто переводи, когда просят.\n"
    "• Писать статьи, посты, письма, поздравления, сценарии, описания — сразу текстом в ответе.\n"
    "• Анализировать и объяснять содержимое файлов и фото: код, документы, таблицы, PDF, архивы.\n"
    "• Находить и исправлять ошибки в коде, объяснять как и почему.\n"
    "• Искать актуальную информацию в интернете, когда это включено в диалоге.\n"
    "• Генерировать НОВЫЕ изображения с нуля по описанию — вызывай функцию generate_image.\n"
    "• Редактировать уже присланное пользователем фото (удалить объект, заменить фон, "
    "улучшить качество, изменить стиль/цвета, добавить объект и т.д.) — вызывай edit_image, "
    "только если фото приложено к текущему сообщению.\n"
    "• Создавать готовые файлы для скачивания (код, документы, таблицы, презентации, "
    "исправленные версии загруженных файлов) — вызывай функцию create_file, когда "
    "пользователь хочет получить именно файл, а не текст в чате.\n\n"
    "Примеры того, как понимать задачу без уточнений: «переведи это на английский» — "
    "просто переведи; «нарисуй красного дракона» — вызови generate_image; «исправь этот "
    "код» — если нужен готовый файл, вызови create_file, если пользователь просто просит "
    "показать исправление — покажи в ответе кодом; «удали человека с фото» — вызови "
    "edit_image; «создай презентацию» — вызови create_file с file_type=pptx; «проанализируй "
    "PDF» — просто опиши/разбери содержимое приложенного файла обычным ответом.\n\n"
    "Отвечай на языке пользователя, будь конкретным и по делу. "
    "Форматируй ответ для Telegram: используй <b>жирный</b> для важного, "
    "<code>код</code> для кода, короткие абзацы, эмодзи по делу — не перегружай. "
    "Если какая-то функция недоступна — вежливо объясни это без технического жаргона."
)

def _owner_system_prompt() -> str:
    who = f"@{SUPER_OWNER_USERNAME}" if SUPER_OWNER_USERNAME else "создателем"
    return (
        BASE_SYSTEM_PROMPT +
        f"\n\nВАЖНО: Ты сейчас общаешься с {who} — создателем и владельцем "
        "проекта Neuravix AI. Обращайся к нему уважительно, ты рад помочь своему "
        "создателю. Поскольку это создатель — можешь быть более открытым и "
        "детальным в ответах."
    )


TITLE_PROMPT = (
    "На основе диалога придумай короткое название чата (3-5 слов, без кавычек и точки в конце), "
    "отражающее суть темы. Ответь только названием, без пояснений."
)


def get_system_prompt(user_id: int) -> str:
    if SUPER_OWNER_ID and user_id == SUPER_OWNER_ID:
        return _owner_system_prompt()
    return BASE_SYSTEM_PROMPT


# ── Форматирование ошибок ────────────────────────────────────────────────────

def _format_error(e: Exception) -> str:
    if isinstance(e, (genai_errors.ClientError, genai_errors.ServerError, asyncio.TimeoutError)):
        logger.warning("Ошибка Gemini API: %r", e)
    else:
        # Незнакомая/непредвиденная ошибка — пишем полный traceback,
        # чтобы причину было видно в Railway Logs, а не только в сообщении пользователю.
        logger.error("Непредвиденная ошибка при обращении к Gemini: %r", e, exc_info=True)
    if isinstance(e, genai_errors.ClientError):
        code = getattr(e, "code", None)
        msg = str(e).lower()
        if code == 429 or "resource_exhausted" in msg or "quota" in msg:
            return (
                "⏳ <b>Слишком много запросов.</b>\n\n"
                "Подожди немного и попробуй снова."
            )
        if code == 403 or "permission" in msg or "api key" in msg:
            contact = f"Сообщи администратору @{SUPPORT_USERNAME}." if SUPPORT_USERNAME else "Сообщи об этом администратору бота."
            return (
                f"🔑 <b>Проблема с доступом к ИИ.</b>\n\n"
                f"{contact}"
            )
        if code == 400 and ("safety" in msg or "blocked" in msg):
            return "🚫 <b>Запрос заблокирован фильтром безопасности.</b>\n\nПопробуй перефразировать."
        if code == 400:
            return "⚠️ <b>Не удалось обработать запрос.</b>\n\nПопробуй сформулировать иначе."
        return "❌ <b>Сервис ИИ временно недоступен.</b>\n\nПопробуй снова через несколько секунд."
    if isinstance(e, genai_errors.ServerError):
        return "🌐 <b>Сервис ИИ перегружен.</b>\n\nПодожди немного и попробуй снова."
    if isinstance(e, asyncio.TimeoutError):
        return "⏱ <b>Превышено время ожидания.</b>\n\nПопробуй ещё раз."
    msg = str(e).lower()
    if "context" in msg or "token" in msg:
        return "📝 <b>Слишком длинный диалог.</b>\n\nНачни новый чат — нажми ➕."
    return "❌ <b>Что-то пошло не так.</b>\n\nПопробуй ещё раз."


class AIError(Exception):
    """Уже отформатированное для пользователя сообщение об ошибке."""
    pass


class ToolCall:
    """Модель решила вызвать функцию (сгенерировать/отредактировать картинку,
    создать файл) вместо обычного текстового ответа."""
    def __init__(self, name: str, args: dict):
        self.name = name
        self.args = args or {}


# ── Инструменты (function calling) ──────────────────────────────────────────
# Позволяют нейросети САМОЙ решать, когда нужно нарисовать/отредактировать
# изображение или создать файл — без отдельных кнопок и режимов.
# ВАЖНО: Gemini 2.x не поддерживает одновременно google_search и
# function_declarations в одном запросе, поэтому инструменты включаются
# только когда веб-поиск выключен (см. _tools_config ниже).

TOOLS = [
    types.Tool(function_declarations=[
        types.FunctionDeclaration(
            name="generate_image",
            description=(
                "Сгенерировать новое изображение с нуля по текстовому описанию. "
                "Используй, когда пользователь просит нарисовать, создать, "
                "сгенерировать картинку/фото/арт/иллюстрацию с нуля (не редактирование "
                "уже существующего фото)."
            ),
            parameters=types.Schema(
                type=types.Type.OBJECT,
                properties={
                    "prompt": types.Schema(
                        type=types.Type.STRING,
                        description=(
                            "Подробное описание изображения для генерации: сюжет, стиль, "
                            "цвета, освещение, ракурс. Лучше на английском для максимального "
                            "качества, но допустим и русский."
                        ),
                    ),
                },
                required=["prompt"],
            ),
        ),
        types.FunctionDeclaration(
            name="edit_image",
            description=(
                "Отредактировать фотографию, которую пользователь прислал В ЭТОМ ЖЕ "
                "сообщении. Используй для: удаления объектов/людей с фото, замены фона, "
                "улучшения качества, изменения стиля, добавления объектов, изменения "
                "цветов и любых других правок существующего изображения. НЕ вызывай эту "
                "функцию, если в текущем сообщении нет прикреплённого фото."
            ),
            parameters=types.Schema(
                type=types.Type.OBJECT,
                properties={
                    "instruction": types.Schema(
                        type=types.Type.STRING,
                        description="Чёткое описание того, что нужно изменить на фото.",
                    ),
                },
                required=["instruction"],
            ),
        ),
        types.FunctionDeclaration(
            name="create_file",
            description=(
                "Создать готовый файл для скачивания: код, документ, статью, презентацию, "
                "таблицу, исправленную версию загруженного файла и т.п. Используй, когда "
                "пользователь просит создать/сохранить/отредактировать файл, документ, код, "
                "презентацию или таблицу и ожидает получить готовый файл, а не просто текст "
                "в чате (например: 'пришли файлом', 'сделай docx', 'создай презентацию', "
                "'исправь код и пришли готовый файл')."
            ),
            parameters=types.Schema(
                type=types.Type.OBJECT,
                properties={
                    "filename": types.Schema(
                        type=types.Type.STRING,
                        description="Имя файла с расширением, например script.py, статья.docx, план.xlsx, презентация.pptx",
                    ),
                    "file_type": types.Schema(
                        type=types.Type.STRING,
                        enum=["text", "docx", "xlsx", "pptx"],
                        description=(
                            "text — код/txt/md/json/html и т.п. (содержимое как есть); "
                            "docx — документ Word (используй заголовки строками, начинающимися с '# '); "
                            "xlsx — таблица (строки через перенос строки, ячейки через ' | '); "
                            "pptx — презентация (слайды разделяй строкой '---', первая строка слайда — заголовок)."
                        ),
                    ),
                    "content": types.Schema(
                        type=types.Type.STRING,
                        description="Полное содержимое файла целиком (не сокращай и не описывай — выдай готовый результат).",
                    ),
                },
                required=["filename", "file_type", "content"],
            ),
        ),
    ]),
]


def _to_history(messages: list[dict]) -> list[types.Content]:
    contents = []
    for m in messages:
        role = "model" if m["role"] in ("model", "assistant") else "user"
        parts = []
        if m.get("content"):
            parts.append(types.Part.from_text(text=m["content"]))
        if parts:
            contents.append(types.Content(role=role, parts=parts))
    return contents


# ── Извлечение текста из файлов ─────────────────────────────────────────────

def extract_text_from_file(filename: str, data: bytes) -> str | None:
    name = filename.lower()
    try:
        if name.endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs if p.text)
        if name.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"--- Лист: {ws.title} ---")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        lines.append(" | ".join("" if c is None else str(c) for c in row))
            return "\n".join(lines[:2000])
        if name.endswith(".zip"):
            zf = zipfile.ZipFile(io.BytesIO(data))
            names = zf.namelist()
            return "Содержимое архива (" + str(len(names)) + " файлов):\n" + "\n".join(names[:300])
        if name.endswith((
            ".py", ".js", ".ts", ".html", ".css", ".json", ".java", ".kt", ".txt", ".md",
            ".yml", ".yaml", ".xml", ".cfg", ".ini", ".sh", ".c", ".cpp", ".h", ".cs",
            ".mcfunction", ".properties", ".toml",
        )):
            import chardet
            enc = chardet.detect(data).get("encoding") or "utf-8"
            return data.decode(enc, errors="replace")[:30000]
    except Exception:
        return None
    return None


GEMINI_NATIVE_MIME_PREFIXES = ("image/", "audio/", "video/")
GEMINI_NATIVE_EXACT = ("application/pdf",)


def guess_mime(filename: str) -> str:
    import mimetypes
    mime, _ = mimetypes.guess_type(filename)
    return mime or "application/octet-stream"


# ── Потоковый чат ────────────────────────────────────────────────────────────

def _is_overloaded(e: Exception) -> bool:
    if isinstance(e, genai_errors.ServerError):
        return True
    if isinstance(e, genai_errors.ClientError):
        code = getattr(e, "code", None)
        msg = str(e).lower()
        return code == 429 or "resource_exhausted" in msg or "quota" in msg
    return False


async def stream_chat(
    history: list[dict],
    model: str,
    max_tokens: int = 4096,
    use_search: bool = False,
    file_bytes: bytes | None = None,
    file_name: str | None = None,
    file_mime: str | None = None,
    user_id: int = 0,
    enable_tools: bool = True,
):
    """
    Асинхронный генератор, отдающий кусочки текста по мере генерации.
    history — последнее сообщение уже включено как {"role": "user", "content": ...}.
    Бросает AIError с готовым HTML-текстом при сбое.
    Может отдать объект ToolCall вместо текста, если модель решила вызвать
    инструмент (сгенерировать/отредактировать изображение, создать файл) —
    в таком случае это последний элемент генератора.
    """
    client = get_client()
    contents = _to_history(history)
    if not contents:
        raise AIError("❌ <b>Пустой запрос.</b>\n\nНапиши что-нибудь!")

    system_prompt = get_system_prompt(user_id)
    has_image_attachment = bool(file_bytes) and (file_mime or "").startswith("image/")
    if has_image_attachment and enable_tools and not use_search:
        system_prompt += (
            "\n\nК текущему сообщению пользователя приложено изображение. Если "
            "пользователь просит его отредактировать (удалить/заменить/улучшить/"
            "изменить что-то на фото) — вызови функцию edit_image. Если он просто "
            "просит описать/проанализировать фото — отвечай обычным текстом."
        )

    # Прикрепляем файл к последнему сообщению пользователя
    if file_bytes and contents:
        mime = file_mime or guess_mime(file_name or "")
        last = contents[-1]
        if mime.startswith(GEMINI_NATIVE_MIME_PREFIXES) or mime in GEMINI_NATIVE_EXACT:
            last.parts.append(types.Part.from_bytes(data=file_bytes, mime_type=mime))
        else:
            extracted = extract_text_from_file(file_name or "file", file_bytes)
            if extracted:
                last.parts.append(types.Part.from_text(
                    text=f"\n\n[Содержимое файла «{file_name}»]:\n{extracted[:15000]}"
                ))
            else:
                last.parts.append(types.Part.from_text(
                    text=f"\n\n[Файл «{file_name}» — формат не поддерживается для чтения.]"
                ))

    def _tools(with_search: bool, with_tools: bool):
        # Gemini 2.x не умеет одновременно google_search и function_declarations —
        # поиск в приоритете, если он явно включён пользователем в диалоге.
        if with_search:
            return [types.Tool(google_search=types.GoogleSearch())]
        if with_tools:
            return TOOLS
        return None

    def _config(with_search: bool, with_tools: bool):
        return types.GenerateContentConfig(
            system_instruction=system_prompt,
            max_output_tokens=max_tokens,
            tools=_tools(with_search, with_tools),
        )

    attempts = [(model, use_search)]
    if use_search:
        attempts.append((model, False))
    if CHAT_FALLBACK_MODEL and CHAT_FALLBACK_MODEL != model:
        attempts.append((CHAT_FALLBACK_MODEL, False))

    last_error: Exception | None = None
    for attempt_index, (try_model, try_search) in enumerate(attempts):
        is_last = attempt_index == len(attempts) - 1
        yielded_any = False
        try:
            async def _run_stream():
                stream = await client.aio.models.generate_content_stream(
                    model=try_model,
                    contents=contents,
                    config=_config(try_search, enable_tools and not try_search),
                )
                pieces = []
                async for chunk in stream:
                    text = None
                    try:
                        text = chunk.text
                    except Exception:
                        text = None
                    if text:
                        pieces.append(text)
                        yield text
                        continue
                    # Проверяем, не решила ли модель вызвать инструмент
                    try:
                        cand = chunk.candidates[0] if chunk.candidates else None
                        if cand and cand.content and cand.content.parts:
                            for part in cand.content.parts:
                                fc = getattr(part, "function_call", None)
                                if fc and getattr(fc, "name", None):
                                    args = dict(fc.args) if fc.args else {}
                                    yield ToolCall(fc.name, args)
                                    return
                    except Exception:
                        pass
            # Обёртываем в таймаут
            gen = _run_stream()
            async def _with_timeout():
                try:
                    async with asyncio.timeout(STREAM_TIMEOUT):
                        async for piece in gen:
                            yield piece
                except asyncio.TimeoutError:
                    raise asyncio.TimeoutError()
            async for piece in _with_timeout():
                yielded_any = True
                yield piece
            return
        except asyncio.TimeoutError as e:
            last_error = e
            # Если часть ответа уже была отдана пользователю — не переключаемся
            # на другую модель заново (иначе текст задвоится/склеится криво).
            if yielded_any:
                raise AIError(_format_error(e)) from e
            if not is_last:
                await asyncio.sleep(1.0)
                continue
            raise AIError(_format_error(e)) from e
        except (genai_errors.ClientError, genai_errors.ServerError) as e:
            last_error = e
            if yielded_any:
                raise AIError(_format_error(e)) from e
            if _is_overloaded(e) and not is_last:
                await asyncio.sleep(1.5)
                continue
            raise AIError(_format_error(e)) from e
        except Exception as e:
            last_error = e
            if yielded_any:
                raise AIError(_format_error(e)) from e
            if not is_last:
                await asyncio.sleep(1.0)
                continue
            raise AIError(_format_error(e)) from e

    if last_error:
        raise AIError(_format_error(last_error)) from last_error


async def generate_title(history: list[dict]) -> str | None:
    try:
        client = get_client()
        convo_text = "\n".join(f"{m['role']}: {m['content'][:300]}" for m in history[:6])
        resp = await asyncio.wait_for(
            client.aio.models.generate_content(
                model=UTILITY_MODEL,
                contents=[types.Content(role="user", parts=[types.Part.from_text(
                    text=f"{TITLE_PROMPT}\n\n{convo_text}"
                )])],
                config=types.GenerateContentConfig(max_output_tokens=30),
            ),
            timeout=10.0
        )
        title = (resp.text or "").strip().strip('"').strip("«»").strip()
        return title[:60] if title else None
    except Exception:
        return None


async def edit_image(image_bytes: bytes, image_mime: str, instruction: str) -> bytes:
    """
    Редактирует уже существующее изображение по текстовой инструкции
    (удалить объект, заменить фон, изменить стиль/цвета и т.д.).
    Использует те же мультимодальные модели, что и генерация, отправляя
    исходное фото + инструкцию.
    """
    client = get_client()
    image_mime = image_mime or "image/jpeg"

    async def _try(model_name: str) -> bytes | None:
        try:
            resp = await asyncio.wait_for(
                client.aio.models.generate_content(
                    model=model_name,
                    contents=[types.Content(role="user", parts=[
                        types.Part.from_bytes(data=image_bytes, mime_type=image_mime),
                        types.Part.from_text(text=instruction),
                    ])],
                    config=types.GenerateContentConfig(response_modalities=["TEXT", "IMAGE"]),
                ),
                timeout=40.0,
            )
            if resp.candidates:
                for cand in resp.candidates:
                    if cand.content and cand.content.parts:
                        for part in cand.content.parts:
                            inline = getattr(part, "inline_data", None)
                            if inline and getattr(inline, "data", None):
                                return inline.data
        except Exception as e:
            logger.warning("Ошибка edit_image (%s): %r", model_name, e)
        return None

    for model_name in IMAGE_GEN_MODELS:
        data = await _try(model_name)
        if data:
            return data

    raise AIError(
        "🖼 <b>Не удалось отредактировать изображение.</b>\n\n"
        "Попробуй переформулировать, что нужно изменить, или повтори чуть позже."
    )


# ── Генерация изображений ────────────────────────────────────────────────────

async def generate_image(prompt: str) -> bytes:
    """
    Генерирует изображение по текстовому описанию.
    Использует несколько моделей с fallback.
    """
    client = get_client()

    # Метод 1: generate_content с модальностями TEXT+IMAGE
    async def _try_generate_content(model_name: str) -> bytes | None:
        for modalities in [["TEXT", "IMAGE"], ["IMAGE"]]:
            try:
                resp = await asyncio.wait_for(
                    client.aio.models.generate_content(
                        model=model_name,
                        contents=prompt,
                        config=types.GenerateContentConfig(
                            response_modalities=modalities,
                        ),
                    ),
                    timeout=30.0,
                )
                if resp.candidates:
                    for cand in resp.candidates:
                        if cand.content and cand.content.parts:
                            for part in cand.content.parts:
                                inline = getattr(part, "inline_data", None)
                                if inline and getattr(inline, "data", None):
                                    return inline.data
            except asyncio.TimeoutError:
                pass
            except Exception:
                pass
        return None

    # Метод 2: Imagen API
    async def _try_imagen(model_name: str) -> bytes | None:
        try:
            resp = await asyncio.wait_for(
                client.aio.models.generate_images(
                    model=model_name,
                    prompt=prompt,
                    config=types.GenerateImagesConfig(
                        number_of_images=1,
                        safety_filter_level="block_only_high",
                        person_generation="allow_adult",
                    ),
                ),
                timeout=30.0,
            )
            if resp.generated_images:
                img = resp.generated_images[0]
                if img.image and getattr(img.image, "image_bytes", None):
                    return img.image.image_bytes
        except Exception:
            pass
        return None

    # Пробуем все методы по порядку
    for model in IMAGE_GEN_MODELS:
        data = await _try_generate_content(model)
        if data:
            return data

    # Пробуем Imagen
    for imagen_model in ["imagen-3.0-generate-001", "imagen-3.0-fast-generate-001"]:
        data = await _try_imagen(imagen_model)
        if data:
            return data

    raise AIError(
        "🖼 <b>Генерация изображений временно недоступна.</b>\n\n"
        "Возможные причины: ограничения API, перегрузка сервиса. "
        "Попробуй позже или обратись к администратору."
    )


# ── Перевод / определение языка ──────────────────────────────────────────────

async def translate_text(text: str, target_language: str) -> str:
    client = get_client()
    try:
        resp = await asyncio.wait_for(
            client.aio.models.generate_content(
                model=UTILITY_MODEL,
                contents=[types.Content(role="user", parts=[types.Part.from_text(text=text)])],
                config=types.GenerateContentConfig(
                    system_instruction=(
                        f"Ты профессиональный переводчик. Переведи текст на {target_language}. "
                        "Выдай только перевод, без пояснений и комментариев."
                    ),
                    max_output_tokens=2000,
                ),
            ),
            timeout=20.0,
        )
        return resp.text or "Не удалось перевести текст."
    except (genai_errors.ClientError, genai_errors.ServerError) as e:
        return _format_error(e)
    except Exception as e:
        return _format_error(e)


async def detect_language(text: str) -> str:
    client = get_client()
    try:
        resp = await asyncio.wait_for(
            client.aio.models.generate_content(
                model=UTILITY_MODEL,
                contents=[types.Content(role="user", parts=[types.Part.from_text(text=text[:200])])],
                config=types.GenerateContentConfig(
                    system_instruction=(
                        "Определи язык текста и ответь только его названием на русском. "
                        "Например: Русский, Английский, Немецкий."
                    ),
                    max_output_tokens=20,
                ),
            ),
            timeout=10.0,
        )
        return resp.text or "Неизвестный"
    except Exception:
        return "Неизвестный"


# ── Красивый текст ────────────────────────────────────────────────────────────

async def write_text(text_type: str, topic: str, model: str = None, max_tokens: int = 2000) -> str:
    client = get_client()
    type_prompts = {
        "article": "Напиши развёрнутую, интересную и структурированную статью на тему",
        "post": "Напиши цепляющий, живой пост для социальных сетей на тему",
        "congrats": "Напиши красивое, тёплое, искреннее поздравление на тему",
        "letter": "Напиши вежливое, грамотное и профессиональное письмо на тему",
        "description": "Напиши яркое, продающее описание для",
        "scenario": "Напиши увлекательный, детальный сценарий на тему",
        "free": "Напиши красивый, интересный, структурированный текст на тему",
    }
    prompt = type_prompts.get(text_type, "Напиши текст на тему")
    try:
        resp = await asyncio.wait_for(
            client.aio.models.generate_content(
                model=model or UTILITY_MODEL,
                contents=[types.Content(role="user", parts=[types.Part.from_text(
                    text=f"{prompt}: {topic}"
                )])],
                config=types.GenerateContentConfig(
                    system_instruction=(
                        "Ты профессиональный копирайтер и писатель. Пишешь красиво, грамотно, интересно. "
                        "Используй структуру с заголовками и абзацами. "
                        "Форматируй для Telegram: <b>жирный</b> для заголовков, абзацы через пустую строку."
                    ),
                    max_output_tokens=max_tokens,
                ),
            ),
            timeout=25.0,
        )
        return resp.text or "Не удалось сгенерировать текст."
    except (genai_errors.ClientError, genai_errors.ServerError) as e:
        return _format_error(e)
    except Exception as e:
        return _format_error(e)
